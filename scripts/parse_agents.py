"""
parse_agents.py
Converts all 3 Precise AI Agent Design xlsx files into JSON.
Handles two distinct schemas:
  - Scored agents (Cardiology + Kidney): Domain/Question/ConfirmingAnswer/Pts/Cost
  - Checklist agents (Neurology): Domain/Question only, no scoring
"""

import json
import re
import sys
from pathlib import Path

import openpyxl

ROOT = Path(__file__).parent.parent
AGENT_SOURCES = {
    "cardiology": ROOT / "agents/cardiology/Cardiology_AI_Agents_v060526.xlsx",
    "kidney":     ROOT / "agents/kidney/Kidney_AI_Agents_v060526.xlsx",
    "neurology":  ROOT / "agents/neurology/Neurology_AI_Agents_v052526.xlsx",
}
OUT_ROOT = ROOT / "src/data/agents"

# Map sheet name → output slug + condition metadata
SHEET_MAP = {
    # Cardiology diagnosis
    "1 BAV":               ("bav",          "diagnosis", "Bicuspid Aortic Valve",            1,  "cardiology"),
    "25 HCM":              ("hcm",          "diagnosis", "Hypertrophic Cardiomyopathy",       25, "cardiology"),
    "36 DCM":              ("dcm",          "diagnosis", "Dilated Cardiomyopathy",             36, "cardiology"),
    "40 Kawasaki":         ("kawasaki",     "diagnosis", "Kawasaki Disease",                  40, "cardiology"),
    "56 PAH":              ("pah",          "diagnosis", "Pulmonary Arterial Hypertension",   56, "cardiology"),
    "59 LQTS":             ("lqts",         "diagnosis", "Long QT Syndrome",                  59, "cardiology"),
    "80 Pulmonary Valve RVOT": ("rvot",     "diagnosis", "Pulmonary Valve / RVOT Disease",    80, "cardiology"),
    "Maternal":            ("maternal",     "diagnosis", "Maternal Cardiac Risk",              0,  "cardiology"),
    # Cardiology progression
    "BAV Prog":            ("bav",          "progression", "Bicuspid Aortic Valve",           1,  "cardiology"),
    "HCM Prog":            ("hcm",          "progression", "Hypertrophic Cardiomyopathy",     25, "cardiology"),
    "DCM Prog":            ("dcm",          "progression", "Dilated Cardiomyopathy",           36, "cardiology"),
    "Kawasaki Prog":       ("kawasaki",     "progression", "Kawasaki Disease",                40, "cardiology"),
    "PAH Prog":            ("pah",          "progression", "Pulmonary Arterial Hypertension", 56, "cardiology"),
    "LQTS Prog":           ("lqts",         "progression", "Long QT Syndrome",                59, "cardiology"),
    "RVOT Prog":           ("rvot",         "progression", "Pulmonary Valve / RVOT Disease",  80, "cardiology"),
    # Kidney diagnosis
    "#2 FSGS Primary":     ("fsgs-primary",   "diagnosis", "FSGS — Primary/Idiopathic",       2,  "kidney"),
    "#3 FSGS Secondary":   ("fsgs-secondary", "diagnosis", "FSGS — Secondary",                3,  "kidney"),
    "#6 IgAN":             ("igan",           "diagnosis", "IgA Nephropathy",                  6,  "kidney"),
    "#110 aHUS":           ("ahus",           "diagnosis", "Atypical Hemolytic Uremic Syndrome", 110, "kidney"),
    "#4 APOL1":            ("apol1",          "diagnosis", "APOL1-Mediated Nephropathy",       4,  "kidney"),
    "#10 C3G":             ("c3g",            "diagnosis", "C3 Glomerulopathy",               10, "kidney"),
    # Kidney progression
    "FSGS Primary Prog":   ("fsgs-primary",   "progression", "FSGS — Primary/Idiopathic",    2,  "kidney"),
    "FSGS Secondary Prog": ("fsgs-secondary", "progression", "FSGS — Secondary",             3,  "kidney"),
    "IgAN Prog":           ("igan",           "progression", "IgA Nephropathy",               6,  "kidney"),
    "aHUS Prog":           ("ahus",           "progression", "Atypical Hemolytic Uremic Syndrome", 110, "kidney"),
    "APOL1 Prog":          ("apol1",          "progression", "APOL1-Mediated Nephropathy",    4,  "kidney"),
    "C3G Prog":            ("c3g",            "progression", "C3 Glomerulopathy",            10,  "kidney"),
    # Neurology (checklist format)
    "#1 Dravet":           ("dravet",        "diagnosis", "Dravet Syndrome",                  1,  "neurology"),
    "#20 TSC":             ("tsc",           "diagnosis", "Tuberous Sclerosis Complex",       20,  "neurology"),
    "#33 Rett":            ("rett",          "diagnosis", "Rett Syndrome",                   33,  "neurology"),
    "#35 Angelman":        ("angelman",      "diagnosis", "Angelman Syndrome",               35,  "neurology"),
    "#54 BEXS":            ("bexs",          "diagnosis", "BEXS Syndrome",                   54,  "neurology"),
    "#55 SMA":             ("sma",           "diagnosis", "Spinal Muscular Atrophy",         55,  "neurology"),
    "#56 DMD":             ("dmd",           "diagnosis", "Duchenne Muscular Dystrophy",     56,  "neurology"),
    "#86 MPS II Hunter":   ("mps2-hunter",   "diagnosis", "MPS II (Hunter Syndrome)",        86,  "neurology"),
    "#87 MPS IIIA Sanfilippo": ("mps3a-sanfilippo", "diagnosis", "MPS IIIA (Sanfilippo)",   87,  "neurology"),
}


def cell_str(cell) -> str:
    """Return stripped string value of a cell, or '' if None."""
    if cell is None or cell.value is None:
        return ""
    return str(cell.value).strip()


def cell_num(cell):
    """Return numeric value of cell, or None."""
    if cell is None or cell.value is None:
        return None
    try:
        return float(cell.value)
    except (TypeError, ValueError):
        return None


def is_confidence_row(row_cells) -> bool:
    """Detect 🔴🟡🟢 confidence interpretation rows."""
    text = cell_str(row_cells[0])
    return any(sym in text for sym in ["🔴", "🟡", "🟢", "High confidence", "Moderate confidence", "Low confidence",
                                        "Fast Progression", "Moderate Progression", "Slow Progression",
                                        "CONFIDENCE SCORE", "PROGRESSION SCORE"])


def parse_confidence_rows(conf_rows: list[str]) -> dict:
    """Parse up to 3 confidence rows into structured thresholds."""
    thresholds = {}
    for text in conf_rows:
        if not text:
            continue
        # Detect tier by emoji or keyword
        if "🔴" in text or "High confidence" in text or "Fast Progression" in text:
            tier = "high"
        elif "🟡" in text or "Moderate confidence" in text or "Moderate Progression" in text:
            tier = "moderate"
        elif "🟢" in text or "Low confidence" in text or "Slow Progression" in text:
            tier = "low"
        else:
            continue

        # Extract numeric range if present (e.g. "80–100" or "66–100" or "<50" or "0–20")
        range_match = re.search(r'(\d+)\s*[–-]\s*(\d+)', text)
        lt_match = re.search(r'<\s*(\d+)', text)
        gte_match = re.search(r'[≥>=]\s*(\d+)', text)

        entry = {"label": text.split("—")[0].strip() if "—" in text else text[:80]}
        # Action text = everything after the first em-dash
        if "—" in text:
            entry["action"] = text.split("—", 1)[1].strip()
        else:
            entry["action"] = text.strip()

        if range_match:
            entry["min"] = int(range_match.group(1))
            entry["max"] = int(range_match.group(2))
        elif lt_match:
            entry["max"] = int(lt_match.group(1)) - 1
            entry["min"] = 0
        elif gte_match:
            entry["min"] = int(gte_match.group(1))
            entry["max"] = 100

        thresholds[tier] = entry
    return thresholds


def parse_scored_sheet(ws) -> dict:
    """
    Parse a Cardiology or Kidney scored agent sheet.
    Returns a ScoredAgent-shaped dict.
    """
    rows = list(ws.iter_rows())
    if not rows:
        return {}

    # Row 0: condition title (col 0)
    condition_title = cell_str(rows[0][0]) if rows else ""

    # Row 1: preamble / instruction text (col 0)
    preamble = cell_str(rows[1][0]) if len(rows) > 1 else ""

    # Row 2 onwards: actual data (header row at index 2 is Domain/Question/...)
    # We skip it since we use positional columns
    sections = []
    current_section = None
    confidence_rows_text = []
    total_points = None
    guidelines_citation = ""

    for row in rows[3:]:  # data starts at row index 3 (0-based)
        col0 = cell_str(row[0])  # Domain / Section heading
        col1 = cell_str(row[1])  # Question
        col2 = cell_str(row[2])  # Confirming Answer
        col3 = cell_num(row[3])  # Points
        col4 = cell_str(row[4])  # Cost

        # Skip completely empty rows
        if not col0 and not col1 and col3 is None:
            continue

        # TOTAL row
        if col0.upper() == "TOTAL" or col1.upper() == "TOTAL":
            if col3 is not None:
                total_points = int(col3)
            continue

        # Confidence / interpretation rows — col0 has emoji+range, col1 has action text
        if is_confidence_row(row):
            # Combine: "🔴  80–100" + " " + "High confidence — ..."
            conf_text = (col0 + (" " + col1 if col1 else "")).strip()
            if conf_text:
                confidence_rows_text.append(conf_text)
            continue

        # Section header: col0 has text, col1 is blank, col3 is section total
        if col0 and not col1 and col3 is not None:
            current_section = {
                "heading": col0,
                "sectionPoints": int(col3),
                "questions": []
            }
            sections.append(current_section)
            continue

        # Question row: col0=domain, col1=question, col3=points
        if col1 and col3 is not None and col3 > 0:
            if current_section is None:
                # Create an implicit section if we encounter a question without one
                current_section = {"heading": col0 or "General", "sectionPoints": 0, "questions": []}
                sections.append(current_section)

            question = {
                "domain": col0 or (current_section["heading"] if current_section else ""),
                "question": col1,
                "confirmingAnswer": col2,
                "points": int(col3),
                "cost": col4 if col4 in ("Small", "Medium", "Large") else "Small",
            }
            current_section["questions"].append(question)
            continue

        # Guidelines citation: long text at end with no points
        if col0 and not col1 and col3 is None and len(col0) > 30:
            if "|" in col0 or "Guidelines" in col0 or "et al" in col0:
                guidelines_citation = col0

    # Parse confidence thresholds
    confidence_thresholds = parse_confidence_rows(confidence_rows_text)

    # Compute total from sections if not explicitly found
    if total_points is None:
        total_points = sum(
            q["points"] for s in sections for q in s["questions"]
        )

    return {
        "type": "scored",
        "totalPoints": total_points,
        "preamble": preamble,
        "sections": sections,
        "confidenceThresholds": confidence_thresholds,
        "guidelinesCitation": guidelines_citation,
    }


def parse_checklist_sheet(ws) -> dict:
    """
    Parse a Neurology checklist sheet (no points, no cost).
    Returns a ChecklistAgent-shaped dict.
    """
    rows = list(ws.iter_rows())
    if not rows:
        return {}

    preamble = cell_str(rows[0][0]) if rows else ""

    sections = []
    current_section = None
    confidence_text = ""

    for row in rows[2:]:
        col0 = cell_str(row[0])
        col1 = cell_str(row[1])

        if not col0 and not col1:
            continue

        # Confidence block (dense multi-line text in col0)
        if is_confidence_row(row):
            confidence_text = col0 or col1
            continue

        # Section header: col0 has text, col1 is blank
        if col0 and not col1:
            current_section = {"heading": col0, "questions": []}
            sections.append(current_section)
            continue

        # Question row
        if col1:
            if current_section is None:
                current_section = {"heading": "General", "questions": []}
                sections.append(current_section)
            current_section["questions"].append({
                "domain": col0 or current_section["heading"],
                "question": col1,
            })

    return {
        "type": "checklist",
        "preamble": preamble,
        "sections": sections,
        "confidenceNote": confidence_text,
    }


def is_neurology_sheet(specialty: str, sheet_name: str) -> bool:
    return specialty == "neurology"


def main():
    # Collect all output per slug: {specialty/slug/agentType → dict}
    # We write one JSON per (slug, agentType) pair
    outputs = {}  # key: (specialty, slug, agent_type) → data dict

    for specialty, xlsx_path in AGENT_SOURCES.items():
        print(f"\n=== Parsing {specialty}: {xlsx_path.name} ===")
        wb = openpyxl.load_workbook(xlsx_path, data_only=True)

        for sheet_name in wb.sheetnames:
            if sheet_name not in SHEET_MAP:
                print(f"  SKIP (not in SHEET_MAP): {sheet_name!r}")
                continue

            slug, agent_type, condition_name, condition_id, spec = SHEET_MAP[sheet_name]
            ws = wb[sheet_name]

            print(f"  Parsing sheet: {sheet_name!r} → {slug}/{agent_type}")

            if is_neurology_sheet(specialty, sheet_name):
                parsed = parse_checklist_sheet(ws)
            else:
                parsed = parse_scored_sheet(ws)

            parsed["conditionId"] = condition_id
            parsed["conditionName"] = condition_name
            parsed["specialty"] = spec
            parsed["agentType"] = agent_type

            key = (spec, slug, agent_type)
            outputs[key] = parsed

    # Write JSON files — one per (slug, agent_type)
    written = 0
    for (specialty, slug, agent_type), data in outputs.items():
        out_dir = OUT_ROOT / specialty
        out_dir.mkdir(parents=True, exist_ok=True)
        filename = f"{slug}-{agent_type}.json"
        out_path = out_dir / filename
        with open(out_path, "w", encoding="utf-8") as f:
            json.dump(data, f, indent=2, ensure_ascii=False)
        q_count = sum(len(s["questions"]) for s in data.get("sections", []))
        print(f"  ✓ {out_path.relative_to(ROOT)}  ({q_count} questions, "
              f"total={data.get('totalPoints', 'N/A')} pts)")
        written += 1

    print(f"\n✅ Done — {written} agent JSON files written to src/data/agents/")


if __name__ == "__main__":
    main()
